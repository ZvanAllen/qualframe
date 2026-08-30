from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import load_workbook


def slugify(text: str) -> str:
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return text


def is_black_fill(rgb: str | None) -> bool:
    return rgb in {None, "00000000", "FF000000"}


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract rightmost colored prompt cells from a workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--sheet", default="Sheet1")
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, data_only=False)
    worksheet = workbook[args.sheet]
    args.output_dir.mkdir(parents=True, exist_ok=True)

    for row_idx in range(2, worksheet.max_row + 1):
        task = worksheet[f"A{row_idx}"].value
        if not task:
            continue

        candidates = []
        for col_idx in range(2, worksheet.max_column + 1):
            cell = worksheet.cell(row_idx, col_idx)
            if not isinstance(cell.value, str):
                continue
            fill_rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor.type == "rgb" else None
            if is_black_fill(fill_rgb):
                continue
            if len(cell.value.strip()) < 40 and "\n" not in cell.value:
                continue
            candidates.append((col_idx, cell.value))

        if not candidates:
            continue

        _, prompt_text = max(candidates, key=lambda item: item[0])
        file_name = f"{row_idx:02d}_{slugify(str(task))}.md"
        (args.output_dir / file_name).write_text(prompt_text.strip() + "\n", encoding="utf-8")
        print(f"Saved {file_name}")


if __name__ == "__main__":
    main()
