from pathlib import Path
import json
import pytest
from openpyxl import Workbook

from qualframe.cfir_review import locate_quote, reference_codes
from qualframe.dashboard import DashboardApp
from qualframe.pipeline.project_runner import ProjectRunner
from qualframe.evidence_review import write_evidence_review_index, load_evidence_review_records, load_evidence_review_actions
from test_evidence_review import _completed_synthetic_run
from test_evidence_review import _make_repo


def test_normalized_locations_preserve_original_unicode_offsets():
    source = '😀 Start. We\u00a0\u00a0can’t attend. End.'
    start, end, kind = locate_quote(source, '“We can\'t attend.”')
    assert kind == 'normalized_quote_candidate'
    assert source[start:end] == 'We\u00a0\u00a0can’t attend.'
    assert locate_quote('repeat repeat', 'repeat')[:2] == (-1, -1)
    assert locate_quote('Changed words.', 'Other words.')[:2] == (-1, -1)


def test_repeated_child_names_have_distinct_code_ids(tmp_path):
    book = Workbook(); sheet = book.active
    for row in [('IV. INDIVIDUALS DOMAIN', None), ('Construct Name', 'Construct Definition'),
                ('A. Leaders', 'Leaders'), ('1. Need', 'Leader needs'),
                ('B. Deliverers', 'Deliverers'), ('1. Need', 'Deliverer needs'),
                ('CFIR OUTCOMES ADDENDUM', None), ('I. ANTECEDENT ASSESSMENTS', 'Group'),
                ('Name', 'Definition'), ('A. Acceptability', 'Acceptability')]:
        sheet.append(row)
    path = tmp_path/'book.xlsx'; book.save(path)
    codes = reference_codes(path)
    children = [c for c in codes if c['raw_label'] == '1. Need']
    assert len(children) == len({c['id'] for c in children}) == 2
    assert codes[-1]['domain'] == 'OUTCOMES ADDENDUM'
    assert not any(c['raw_label'] == 'Name' for c in codes)


def test_note_rows_between_constructs_are_not_codes(tmp_path):
    book = Workbook(); sheet = book.active
    for row in [('III. INNER SETTING DOMAIN', None), ('Construct Name', 'Construct Definition'),
                ('Note: ', 'Constructs A - D apply to the inner setting as a whole.'), ('A. Structure', 'Structure'),
                ('Note:', 'Constructs E - K are specific to the innovation.'), ('E. Tension for Change', 'Tension')]:
        sheet.append(row)
    path = tmp_path/'book.xlsx'; book.save(path)
    assert [c['raw_label'] for c in reference_codes(path)] == ['A. Structure', 'E. Tension for Change']


def test_review_advances_and_rejects_unknown_edit_codes(tmp_path):
    project, run_id, root, review = _completed_synthetic_run(tmp_path)
    index = write_evidence_review_index(project_manifest_path=project, run_id=run_id)
    records = load_evidence_review_records(index)
    app = DashboardApp(ProjectRunner(project))
    form = {'run_id':[run_id], 'item':[records[0].record_id], 'decision':['accept'],
            'reviewer':['Researcher'], 'advance':['yes']}
    app.handle_action('/evidence-action', form)
    assert form['item'] == [records[1].record_id]
    assert len(load_evidence_review_actions(review)) == 1

    form.update(decision=['edit'], edited_start=['0'], edited_end=['2'], edited_code_id=['not-a-code'])
    with pytest.raises(ValueError, match='Choose a code'):
        app.handle_action('/evidence-action', form)
    assert len(load_evidence_review_actions(review)) == 1


def test_reference_import_changes_fingerprint_and_preserves_prior_workbook(tmp_path):
    from io import BytesIO
    from qualframe.models import StudyCreationRequest
    from qualframe.study_setup import StudyWorkspace
    from qualframe.cfir_import import import_reference
    root = _make_repo(tmp_path)
    workspace = StudyWorkspace(root)
    created = workspace.create(StudyCreationRequest(title='CFIR import test', template='cfir',
        data_classification='sensitive', include_synthetic_fixtures=False))
    project = root / created.project_manifest_path
    initial = ProjectRunner(project)
    assert [s.id for s in initial.pipeline.stages if s.enabled] == ['transcript_to_cfir']
    assert initial.project.governance.enforce_review_gates
    def content(definition):
        book = Workbook(); sheet = book.active
        sheet.append(['I. INNOVATION DOMAIN', None]); sheet.append(['A. Source', definition])
        output = BytesIO(); book.save(output); return output.getvalue()
    first_bytes = content('Original definition')
    first = import_reference(workspace, project, 'CFIR.xlsx', first_bytes, 'Researcher', 'First version')
    previous = ProjectRunner(project).configuration_lock.fingerprint_sha256
    second = import_reference(workspace, project, 'CFIR.xlsx', content('Revised definition'), 'Researcher', 'Revision')
    assert first.destination_path != second.destination_path
    assert (root / first.destination_path).read_bytes() == first_bytes
    assert ProjectRunner(project).configuration_lock.fingerprint_sha256 != previous

def test_typography_only_acceptance_and_reviewer_retention(tmp_path):
    from qualframe.evidence_review import append_evidence_review_action, blocking_review_issues
    from qualframe.io.transcript_io import transcript_to_text
    project, run_id, root, review = _completed_synthetic_run(tmp_path)
    index = write_evidence_review_index(project_manifest_path=project, run_id=run_id)
    record = load_evidence_review_records(index)[0]
    source_path = project.parents[2] / record.source_path
    source = transcript_to_text(source_path)
    quoted = record.evidence_span.exact_text.replace(' ', '  ')
    changed = record.model_copy(update={
        'proposed_exact_text': quoted, 'anchoring_method': 'normalized_quote_candidate',
        'validation_issues': ['Formatting differs from the model quotation. Confirm the highlighted source passage with Save changes.']})
    assert blocking_review_issues(changed, source) == []
    action = append_evidence_review_action(record=changed, source_path=source_path,
        review_dir=review, artifact_root=root, decision='accept', reviewer='Researcher', rationale='Confirmed highlighted source.')
    assert action.decision == 'accept'
    invalid = changed.model_copy(update={'validation_issues': changed.validation_issues + ['Unknown construct']})
    with pytest.raises(ValueError, match='Invalid evidence'):
        append_evidence_review_action(record=invalid, source_path=source_path,
            review_dir=review, artifact_root=root, decision='accept', reviewer='Researcher', rationale='Test')
    app = DashboardApp(ProjectRunner(project))
    app.handle_action('/evidence-action', {'run_id':[run_id], 'item':[record.record_id],
        'decision':['accept'], 'reviewer':['Remembered Researcher']})
    restarted = DashboardApp(ProjectRunner(project))
    page = restarted.render_evidence(run_id=run_id)
    assert 'Recording decisions as Remembered Researcher' in page
    assert '<summary>Why the model proposed this' in page
    assert 'id="code-search"' not in page
    assert 'id="edit-controls"' in page


def _outline_book(tmp_path):
    book = Workbook(); sheet = book.active
    for row in [('II. OUTER SETTING DOMAIN', None), ('Construct Name', 'Construct Definition'),
                ('D. Partnerships & Connections', 'Partnerships'),
                ('IV. INDIVIDUALS DOMAIN', None), ('Construct Name', 'Construct Definition'),
                ('C. Opinion Leaders', 'Opinion leaders'), ('1. Need', 'Opinion leader need'),
                ('H. Innovation Deliverers', 'Deliverers'), ('1. Need', 'Deliverer need'), ('4. Motivation', 'Deliverer motivation'),
                ('V. IMPLEMENTATION PROCESS DOMAIN \nProcess', None), ('Construct Name', 'Construct Definition'),
                ('B. Assessing Needs', 'Assessing'), ('1. Innovation Deliverers', 'Assessing deliverers'),
                ('F. Engaging', 'Engaging'), ('1. Innovation Deliverers', 'Engaging deliverers'),
                ('H. Reflecting & Evaluating', 'Reflecting'), ('2. Innovation', 'Reflecting on the innovation'),
                ('CFIR OUTCOMES ADDENDUM', None), ('I. ANTECEDENT ASSESSMENTS', 'Group'), ('Name', 'Definition'),
                ('A. Acceptability', 'Acceptability'),
                ('II. IMPLEMENTATION OUTCOMES', 'Group'), ('Name', 'Definition'),
                ('A. Anticipated Implementation Outcomes', 'Anticipated'), ('1. Adoptability', 'Adoptability'),
                ('B. Actual Implementation Outcomes', 'Actual'), ('1. Adoption', 'Adoption'),
                ('III. INNOVATION OUTCOMES', 'Group'), ('Name', 'Definition'),
                ('A. Innovation Recipient Impact', 'Recipient impact')]:
        sheet.append(row)
    path = tmp_path / 'book.xlsx'; book.save(path)
    return reference_codes(path)


def _label_records(tmp_path, codes, rows_by_domain, policy):
    from types import SimpleNamespace
    from qualframe.cfir_review import cfir_records
    source = tmp_path / 'transcript.txt'
    source.write_text('Participant: We met every week.', encoding='utf-8')
    artifact_path = tmp_path / 'artifact.json'; artifact_path.write_text('{}', encoding='utf-8')
    parsed = {domain: [{'CFIR Construct': label, 'Quote': 'We met every week.'} for label in labels]
              for domain, labels in rows_by_domain.items()}
    artifact = SimpleNamespace(parsed_output=parsed, stage_id='transcript_to_cfir', model='m', provider='p', prompt_sha256='0' * 64)
    manifest = SimpleNamespace(run_id='run', configuration_fingerprint='f')
    return cfir_records(artifact=artifact, artifact_rel='artifact.json', artifact_path=artifact_path, source_path=source,
                        repo_root=tmp_path, manifest=manifest, codes=codes, codebook_hash='h', code_label_policy=policy)


def test_outline_labels_match_parent_child_and_outcomes_groups(tmp_path):
    codes = _outline_book(tmp_path)
    by_id = {c['id']: c for c in codes}
    assert {c['section'] for c in codes if c['domain'] == 'OUTCOMES ADDENDUM'} == {
        'ANTECEDENT ASSESSMENTS', 'IMPLEMENTATION OUTCOMES', 'INNOVATION OUTCOMES'}
    rows = {
        'Outer Setting': ['Partnerships and Connections'],
        'Individuals': ['H. Innovation Deliverers: 4. Motivation', 'C. Opinion Leaders', '1. Need',
                        'H. Innovation Deliverers – 1. Need', 'Z. Somebody Else: 1. Need'],
        'Implementation Process': ['F. Engaging: 1. Innovation Deliverers', 'H. Reflecting and Evaluating: 2. Innovation',
                                   'B. Assessing Needs / 1. Innovation Deliverers'],
        'Outcomes Addendum': ['I.A. Antecedent Assessments: Acceptability',
                              'II.A. Implementation Outcomes: Anticipated Implementation Outcomes',
                              'II.B.1. Actual Implementation Outcomes: Adoption',
                              'III.A. Innovation Outcomes: Innovation Recipient Impact',
                              'II.B.1. Anticipated Implementation Outcomes: Adoption'],
    }
    records = _label_records(tmp_path, codes, rows, 'outline_labels_v3')
    matched = [(by_id[r.proposed_code_id]['parent_label'] or by_id[r.proposed_code_id].get('section'),
                by_id[r.proposed_code_id]['raw_label']) if r.proposed_code_id else None for r in records]
    assert matched == [
        ('', 'D. Partnerships & Connections'),
        ('H. Innovation Deliverers', '4. Motivation'),
        ('', 'C. Opinion Leaders'),
        ('C. Opinion Leaders', '1. Need'),          # unqualified child follows the preceding construct
        ('H. Innovation Deliverers', '1. Need'),
        None,                                       # a qualifier that names no parent is not guessed
        ('F. Engaging', '1. Innovation Deliverers'),
        ('H. Reflecting & Evaluating', '2. Innovation'),
        ('B. Assessing Needs', '1. Innovation Deliverers'),
        ('ANTECEDENT ASSESSMENTS', 'A. Acceptability'),
        ('IMPLEMENTATION OUTCOMES', 'A. Anticipated Implementation Outcomes'),
        ('B. Actual Implementation Outcomes', '1. Adoption'),
        ('INNOVATION OUTCOMES', 'A. Innovation Recipient Impact'),
        None,                                       # Adoption sits under Actual, not Anticipated
    ]
    assert [r.validation_issues for r in records if r.proposed_code_id is None] == [
        ['Construct could not be identified uniquely in the reference codebook.']] * 2
    # The frozen benchmark policies are unchanged: they still miss these forms.
    legacy = _label_records(tmp_path, codes, {'Individuals': ['H. Innovation Deliverers: 4. Motivation']}, 'legacy_v1')
    assert legacy[0].proposed_code_id is None

