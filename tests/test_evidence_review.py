from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from qualframe.dashboard import DashboardApp
from qualframe.directed_analysis import validate_evidence_span_against_source
from qualframe.evidence_review import (
    append_evidence_review_action,
    export_evidence_review_package,
    load_evidence_review_actions,
    load_evidence_review_records,
    write_evidence_review_index,
)
from qualframe.io.files import sha256_for_file, sha256_for_text
from qualframe.io.transcript_io import transcript_to_text
from qualframe.models import EvidenceSpan, StudyCreationRequest
from qualframe.pipeline.project_runner import ProjectRunner
from qualframe.study_setup import StudyWorkspace


REPO_ROOT = Path(__file__).resolve().parents[1]


def _make_repo(tmp_path: Path) -> Path:
    for directory in ("config", "prompts"):
        shutil.copytree(REPO_ROOT / directory, tmp_path / directory)
    shutil.copytree(REPO_ROOT / "demo/fixtures", tmp_path / "demo/fixtures")
    return tmp_path


def _completed_synthetic_run(
    tmp_path: Path, source_format: str = "lf"
) -> tuple[Path, str, Path, Path]:
    repo_root = _make_repo(tmp_path)
    # Convert only disposable fixture copies, before importing immutable sources.
    for source in (repo_root / "demo/fixtures/directed_transcripts").rglob("*.txt"):
        source_text = source.read_text(encoding="utf-8")
        if source_format == "docx":
            from docx import Document

            document = Document()
            for paragraph in source_text.splitlines():
                document.add_paragraph(paragraph)
            document.save(source.with_suffix(".docx"))
            source.unlink()
        else:
            newline = "\r\n" if source_format == "crlf" else "\n"
            source.write_bytes(source_text.replace("\n", newline).encode("utf-8"))
    created = StudyWorkspace(repo_root).create(
        StudyCreationRequest(
            title="Evidence Review Study",
            data_classification="synthetic",
            include_synthetic_fixtures=True,
        )
    )
    project_path = repo_root / created.project_manifest_path
    runner = ProjectRunner(project_path)
    run_id = "evidence-review-run"
    plan = runner.build_run_plan()
    runner.save_run_plan(plan, run_id)
    runner.execute_run_plan(plan, run_id)
    artifact_root = repo_root / "studies" / created.study_id / "artifacts" / run_id
    review_dir = repo_root / "studies" / created.study_id / "review_queue" / run_id
    return project_path, run_id, artifact_root, review_dir


@pytest.mark.parametrize("source_format", ["lf", "crlf", "docx"])
def test_materializes_offset_valid_evidence_from_immutable_sources(
    tmp_path: Path, source_format: str
) -> None:
    project_path, run_id, artifact_root, _review_dir = _completed_synthetic_run(
        tmp_path, source_format
    )

    index_path = write_evidence_review_index(
        project_manifest_path=project_path,
        run_id=run_id,
    )
    records = load_evidence_review_records(index_path)

    manifest = json.loads((artifact_root / "run_manifest.json").read_text())
    assert manifest["qa_failure_count"] == 0
    assert len(records) == 4
    assert {record.assignment.code_id for record in records if record.assignment} == {
        "access_availability",
        "experience_communication",
    }
    assert all(record.evidence_span for record in records)
    assert all(not record.validation_issues for record in records)
    for record in records:
        source_path = project_path.parents[2] / record.source_path
        source_text = transcript_to_text(source_path)
        span = record.evidence_span
        assert span is not None
        assert record.source_sha256 == sha256_for_file(source_path)
        assert span.source_sha256 == sha256_for_text(source_text)
        if source_format != "lf":
            assert span.source_sha256 != record.source_sha256
        validate_evidence_span_against_source(span, source_text)
        assert source_text[span.start_offset : span.end_offset] == span.exact_text
        assert record.context_text and span.exact_text in record.context_text


@pytest.mark.parametrize("source_format", ["lf", "crlf", "docx"])
def test_review_actions_preserve_original_and_export_traceable_package(
    tmp_path: Path, source_format: str
) -> None:
    project_path, run_id, artifact_root, review_dir = _completed_synthetic_run(
        tmp_path, source_format
    )
    records = load_evidence_review_records(
        write_evidence_review_index(project_manifest_path=project_path, run_id=run_id)
    )
    first, second = records[:2]
    accepted = append_evidence_review_action(
        record=first,
        source_path=project_path.parents[2] / first.source_path,
        review_dir=review_dir,
        artifact_root=artifact_root,
        decision="accept",
        reviewer="synthetic reviewer",
        rationale="The exact synthetic span and code are correct.",
    )
    assert second.evidence_span is not None
    edited = append_evidence_review_action(
        record=second,
        source_path=project_path.parents[2] / second.source_path,
        review_dir=review_dir,
        artifact_root=artifact_root,
        decision="edit",
        reviewer="synthetic reviewer",
        rationale="Preserve the span while recording a reviewer-derived assignment.",
        edited_start=second.evidence_span.start_offset,
        edited_end=second.evidence_span.end_offset,
        edited_code_id=second.assignment.code_id if second.assignment else "experience_communication",
    )
    for record in records[2:]:
        append_evidence_review_action(
            record=record,
            source_path=project_path.parents[2] / record.source_path,
            review_dir=review_dir,
            artifact_root=artifact_root,
            decision="accept",
            reviewer="synthetic reviewer",
            rationale="The remaining synthetic evidence record is correct.",
        )

    assert accepted.derived_artifact_path is None
    assert edited.derived_artifact_path
    derived = json.loads((artifact_root / edited.derived_artifact_path).read_text())
    validate_evidence_span_against_source(
        EvidenceSpan.model_validate(derived["evidence_span"]),
        transcript_to_text(project_path.parents[2] / second.source_path),
    )
    assert all(
        sha256_for_file(project_path.parents[2] / record.source_path) == record.source_sha256
        for record in records
    )
    assert len(load_evidence_review_actions(review_dir)) == 4

    written = export_evidence_review_package(
        project_manifest_path=project_path,
        run_id=run_id,
    )
    assert {path.suffix for path in written} == {".json", ".xlsx"}
    package = json.loads(next(path for path in written if path.suffix == ".json").read_text())
    assert package["status"] == "researcher_review_record_not_automatic_findings"
    assert len(package["evidence"]) == 4
    for row in package["evidence"]:
        source_path = project_path.parents[2] / row["Source path"]
        assert row["Source SHA-256"] == sha256_for_file(source_path)
        assert row["Source text SHA-256"] == sha256_for_text(transcript_to_text(source_path))
    workbook = load_workbook(next(path for path in written if path.suffix == ".xlsx"))
    assert workbook.sheetnames == ["Evidence", "Review History", "Methods"]


@pytest.mark.parametrize("decision", ["accept", "edit"])
def test_review_rejects_changed_file_even_when_extracted_text_is_unchanged(
    tmp_path: Path, decision: str
) -> None:
    project_path, run_id, artifact_root, review_dir = _completed_synthetic_run(tmp_path)
    record = load_evidence_review_records(
        write_evidence_review_index(project_manifest_path=project_path, run_id=run_id)
    )[0]
    source_path = project_path.parents[2] / record.source_path
    source_text = transcript_to_text(source_path)
    source_path.write_bytes(source_text.replace("\n", "\r\n").encode("utf-8"))
    assert transcript_to_text(source_path) == source_text
    with pytest.raises(ValueError, match="Source file hash changed"):
        append_evidence_review_action(
            record=record,
            source_path=source_path,
            review_dir=review_dir,
            artifact_root=artifact_root,
            decision=decision,
            reviewer="synthetic reviewer",
            rationale="Changed source bytes must be blocked.",
            edited_start=record.proposed_start_offset,
            edited_end=record.proposed_end_offset,
            edited_code_id="access_availability",
        )
    assert load_evidence_review_actions(review_dir) == []


def test_invalid_evidence_cannot_be_accepted(tmp_path: Path) -> None:
    project_path, run_id, artifact_root, review_dir = _completed_synthetic_run(tmp_path)
    record = load_evidence_review_records(
        write_evidence_review_index(project_manifest_path=project_path, run_id=run_id)
    )[0].model_copy(update={"validation_issues": ["Synthetic mismatch"]})
    with pytest.raises(ValueError, match="cannot be accepted"):
        append_evidence_review_action(
            record=record,
            source_path=project_path.parents[2] / record.source_path,
            review_dir=review_dir,
            artifact_root=artifact_root,
            decision="accept",
            reviewer="synthetic reviewer",
            rationale="This should be blocked.",
        )


def test_dashboard_reviews_all_exact_evidence_and_exports(tmp_path: Path) -> None:
    project_path, run_id, artifact_root, review_dir = _completed_synthetic_run(tmp_path)
    app = DashboardApp(ProjectRunner(project_path))

    selected, message = app.handle_action("/prepare-evidence", {"run_id": [run_id]})
    assert selected == run_id
    assert "Prepared 4" in message
    index_path = artifact_root / "evidence_review/index.json"
    records = load_evidence_review_records(index_path)
    page = app.render_evidence(run_id=run_id, item_id=records[0].record_id)
    assert "Review exact evidence" in page
    assert "<mark>" in page
    assert "not a research finding" in page
    assert records[0].source_sha256 in page

    for record in records:
        _selected, action_message = app.handle_action(
            "/evidence-action",
            {
                "run_id": [run_id],
                "item": [record.record_id],
                "decision": ["accept"],
                "reviewer": ["synthetic browser reviewer"],
                "rationale": ["The source-linked synthetic record is correct."],
                "edited_start": [str(record.proposed_start_offset)],
                "edited_end": [str(record.proposed_end_offset)],
                "edited_code_id": [record.assignment.code_id if record.assignment else ""],
            },
        )
        assert "Recorded evidence decision accept" in action_message

    _selected, export_message = app.handle_action(
        "/export-evidence", {"run_id": [run_id]}
    )
    assert "Exported 2" in export_message
    final_page = app.render_evidence(run_id=run_id, item_id=records[-1].record_id)
    assert "4 / 4" in final_page
    assert "reviewed_evidence_package.json" in final_page
    assert "reviewed_evidence_package.xlsx" in final_page
    assert len(load_evidence_review_actions(review_dir)) == 4
