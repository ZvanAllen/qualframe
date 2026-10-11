from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

from qualframe.qualbench import (
    agreement_statistics,
    build_qualbench,
    krippendorff_alpha_nominal,
    reference_coding_set,
    reference_statistics,
    run_coding_set,
    write_qualbench,
)
from qualframe.reference_coding import (
    build_reference_coding,
    list_references,
    load_reference_coding,
    locate_reference_quote,
    save_reference_coding,
)
from test_evidence_review import _completed_synthetic_run, _make_repo

ROOT = Path(__file__).resolve().parents[1]


def test_agreement_statistics_match_hand_calculation():
    stats = agreement_statistics(20, 5, 10, 15)
    assert stats["n_units"] == 50
    assert stats["percent_agreement"] == pytest.approx(0.7)
    assert stats["cohen_kappa"] == pytest.approx(0.4)
    assert stats["gwet_ac1"] == pytest.approx((0.7 - 0.495) / 0.505)
    assert stats["pabak"] == pytest.approx(0.4)
    assert stats["jaccard"] == pytest.approx(20 / 35)
    assert reference_statistics(20, 5, 10) == pytest.approx(dict(precision=0.8, recall=20 / 30, f1=40 / 55))
    # Every unit present on both sides: agreement is perfect but kappa is undefined.
    assert agreement_statistics(4, 0, 0, 0)["cohen_kappa"] is None
    assert agreement_statistics(0, 0, 0, 0)["percent_agreement"] is None


def test_krippendorff_alpha_nominal():
    assert krippendorff_alpha_nominal([[1, 1], [0, 0], [1, 0], [0, 0]]) == pytest.approx(1 - 7 * 2 / 30)
    assert krippendorff_alpha_nominal([[1, 1, 1], [0, 0, 0]]) == pytest.approx(1.0)
    assert krippendorff_alpha_nominal([[1, 1], [1, 1]]) is None  # no variation: undefined


def _set(set_id, included, excluded=(), kind="model_proposals"):
    codes = [
        dict(id="p", label="Parent", raw_label="A. Parent", parent_label="", domain="FIRST DOMAIN"),
        dict(id="c", label="Parent / Child", raw_label="1. Child", parent_label="A. Parent", domain="FIRST DOMAIN"),
        dict(id="q", label="Other", raw_label="B. Other", parent_label="", domain="OUTCOMES ADDENDUM"),
    ]
    def item(n, source, code, start=None, end=None):
        row = dict(record_id=f"{set_id}-{n}", source=source, code_id=code, anchoring_method="unique_exact_quote")
        if start is not None:
            row.update(start=start, end=end, exact_text="x" * (end - start))
        return row
    return dict(set_id=set_id, set_kind=kind, label=set_id, model=None, provider=None, usage=None,
                codebook=dict(id="book", codes=codes), sources={"t/one.txt": "h1", "t/two.txt": "h2"},
                included=[item(n, *row) for n, row in enumerate(included)],
                excluded=[dict(item(n + 100, *row), issues=["not found"]) for n, row in enumerate(excluded)],
                absences=[])


def test_build_qualbench_presence_passages_levels_and_stability(tmp_path):
    left = _set("left", [("t/one.txt", "p", 0, 10), ("t/one.txt", "c", 20, 30), ("t/two.txt", "q", 0, 5)],
                excluded=[("t/two.txt", "p")])
    right = _set("right", [("t/one.txt", "p", 0, 10), ("t/one.txt", "c", 25, 40)])
    reference = dict(_set("reference:ref", [("t/one.txt", "p", 0, 10), ("t/two.txt", "c", 3, 9)], kind="reference"),
                     statement="Earlier coding, reviewed.")
    report = build_qualbench([left, right], reference=reference, stability_groups=[["left", "right"]])
    rows = {(r["level"], r["scope_kind"], r["scope"], r["left"], r["right"]): r for r in report["presence"]}
    pair = rows[("code", "all", "All", "left", "right")]
    # Units: 2 transcripts x 3 codes. Left: one/p, one/c, two/q, two/p (unlocated still counts).
    assert (pair["both"], pair["left_only"], pair["right_only"], pair["neither"]) == (2, 2, 0, 2)
    assert pair["precision"] is None
    versus = rows[("code", "all", "All", "left", "reference:ref")]
    assert (versus["both"], versus["left_only"], versus["right_only"]) == (1, 3, 1)
    assert versus["recall"] == pytest.approx(0.5)
    assert rows[("code", "domain", "Outcomes Addendum", "left", "right")]["n_units"] == 2
    assert rows[("code", "transcript", "one.txt", "left", "right")]["n_units"] == 3
    # Sub-codes merged into parents: the reference's two/c becomes two/p, which left has.
    merged = rows[("parent code", "all", "All", "left", "reference:ref")]
    assert (merged["both"], merged["left_only"], merged["right_only"]) == (2, 1, 0)
    passages = {(r["left"], r["right"]): r for r in report["passages"] if r["scope_kind"] == "all"}
    lr = passages[("left", "right")]
    per_transcript = [r for r in report["passages"] if r["scope_kind"] == "transcript" and r["right"] == "right"]
    assert [r["scope"] for r in per_transcript if r["left"] == "left"] == ["one.txt", "two.txt"]
    assert lr["matched_iou_quarter"] == 2 and lr["matched_iou_three_quarters"] == 1
    macro = rows[("code", "transcript mean", "Mean of transcripts", "left", "right")]
    by_transcript = [rows[("code", "transcript", name, "left", "right")]["percent_agreement"] for name in ("one.txt", "two.txt")]
    assert macro["percent_agreement"] == pytest.approx(sum(by_transcript) / 2) and macro["both"] is None
    assert (lr["matched_exact"], lr["matched_iou_half"], lr["matched_any_overlap"]) == (1, 1, 2)
    assert lr["coded_text_jaccard"] == pytest.approx((10 + 5) / (10 + 20 + 5))
    [stability_code, stability_parent] = report["stability"]
    assert stability_code["coding_sets"] == 2 and stability_code["krippendorff_alpha"] is not None
    assert len(report["units"]) == 3 * 2 * 3
    assert report["reference"]["statement"] == "Earlier coding, reviewed."
    folder = write_qualbench(report, tmp_path / "bench")
    assert {p.name for p in folder.iterdir()} == {"README.md", "report.json", "coding_sets.csv", "presence.csv",
                                                 "passages.csv", "stability.csv", "units.csv"}
    assert "Earlier coding, reviewed." in (folder / "README.md").read_text(encoding="utf-8")
    with pytest.raises(FileExistsError):
        write_qualbench(report, folder)


def test_qualbench_rejects_mismatched_inputs():
    one = _set("one", [])
    with pytest.raises(ValueError, match="at least two"):
        build_qualbench([one])
    other = _set("two", [])
    other["codebook"] = dict(id="book", codes=other["codebook"]["codes"][:1])
    with pytest.raises(ValueError, match="same codebook"):
        build_qualbench([one, other])
    apart = _set("three", [])
    apart["sources"] = {"t/three.txt": "h3"}
    with pytest.raises(ValueError, match="no transcripts in common"):
        build_qualbench([one, apart])


def test_run_coding_set_reads_a_completed_run(tmp_path):
    project, run_id, run_root, _ = _completed_synthetic_run(tmp_path)
    from qualframe.evidence_review import write_evidence_review_index
    write_evidence_review_index(project_manifest_path=project, run_id=run_id)
    coding = run_coding_set(project.parents[2], run_root)
    assert coding["set_id"] == run_id and coding["set_kind"] == "model_proposals"
    assert coding["usage"]["cost_usd"] == 0 and coding["included"]
    twin = dict(coding, set_id="repeat")
    report = build_qualbench([coding, twin], stability_groups=[[run_id, "repeat"]])
    overall = next(r for r in report["presence"] if r["scope_kind"] == "all" and r["level"] == "code")
    assert overall["percent_agreement"] == 1.0


def test_reference_quotes_shortened_with_an_ellipsis_are_located_by_their_parts():
    source = "We met every week at the clinic. It was hard at first. Later the residents came to rely on it."
    start, end, method = locate_reference_quote(source, "We met every week at the clinic ... came to rely on it.")
    assert method == "ellipsis_parts" and source[start:end].startswith("We met") and source[start:end].endswith("rely on it.")
    assert locate_reference_quote(source, "Something nobody said ... at all, really.")[2] == "unresolved_exact_quote"


def _cfir_study(tmp_path):
    from io import BytesIO
    from qualframe.cfir_import import import_reference
    from qualframe.models import StudyCreationRequest
    from qualframe.study_setup import StudyWorkspace
    root = _make_repo(tmp_path)
    workspace = StudyWorkspace(root)
    created = workspace.create(StudyCreationRequest(title="Reference test", template="cfir",
                                                    data_classification="synthetic", include_synthetic_fixtures=False))
    project = root / created.project_manifest_path
    book = Workbook(); sheet = book.active
    for row in [("III. INNER SETTING DOMAIN", None), ("Construct Name", "Construct Definition"),
                ("A. Structure", "Structure"),
                ("IV. INDIVIDUALS DOMAIN", None), ("Construct Name", "Construct Definition"),
                ("A. Leaders", "Leaders"), ("1. Need", "Leader need"), ("B. Deliverers", "Deliverers"), ("1. Need", "Deliverer need"),
                ("CFIR OUTCOMES ADDENDUM", None), ("I. ANTECEDENT ASSESSMENTS", "Group"), ("Name", "Definition"),
                ("A. Acceptability", "Acceptability")]:
        sheet.append(row)
    buffer = BytesIO(); book.save(buffer)
    import_reference(workspace, project, "CFIR.xlsx", buffer.getvalue(), "Researcher", "Codebook")
    transcript = project.parent / "data/raw/transcripts/ON/1A_Transcript.txt"
    transcript.parent.mkdir(parents=True, exist_ok=True)
    transcript.write_text("Interviewer: How did it go?\nParticipant: Our leaders backed it from the start. "
                          "We met every week in the old library. The mentors loved it, and the residents "
                          "kept coming back for more.\n", encoding="utf-8")
    return root, project, transcript.relative_to(root).as_posix()


def test_reference_coding_import_matches_constructs_locates_quotes_and_keeps_notes(tmp_path):
    root, project, transcript = _cfir_study(tmp_path)
    coding = Workbook()
    individuals = coding.active; individuals.title = "Individuals"
    for row in [("Checklist:",), ("- Full read-through",), (None,),
                ("CFIR Domain", "CFIR Construct", "Quote", "Page Number", "Explanation"),
                ("Individuals", "A. Leaders", "Our leaders backed it from the start.", None, "Facilitating"),
                ("Individuals", "1. Need", "No relevant quote found", None, "None"),
                ("Individuals", "B. Deliverers", None, None, None),
                ("Individuals", "1. Need", "A sentence that is not in the transcript at all.", None, "x"),
                ("Individuals", "Z. Someone Else", "We met every week in the old library.", None, "x")]:
        individuals.append(row)
    inner = coding.create_sheet("Inner Setting")
    inner.append(("CFIR Construct", "Quote", "Page Number", "Explanation", "Notes"))
    inner.append(("A. Structure", "We met every week in the old library.", 2, "Neutral", "ok"))
    outcomes = coding.create_sheet("CIFR Outcomes")
    outcomes.append(("CFIR Domain", "CFIR Construct", "Quote", "Page Number", "Explanation"))
    outcomes.append(("ANTECEDENT ASSESSMENT", "A. Acceptability",
                     "The mentors loved it ... kept coming back for more.", None, "Facilitating"))
    coding.create_sheet("Summary").append(("Totals", 3))
    workbook = tmp_path / "1A coding.xlsx"; coding.save(workbook)

    payload = build_reference_coding(repo_root=root, project_path=project, reference_id="phase-1a",
                                     label="Earlier coding", statement="Model output reviewed by the team.",
                                     files={transcript: workbook}, created_by="Researcher")
    assert payload["counts"] == dict(included=3, excluded=2, absences=1, unmatched_constructs=1, quotations_not_found=1)
    located = {item["construct_label"]: item for item in payload["included"]}
    assert located["A. Acceptability"]["anchoring_method"] == "ellipsis_parts"
    assert located["A. Structure"]["notes"] == {"notes": "ok"}
    assert payload["absences"][0]["construct_label"] == "1. Need"
    missing = next(i for i in payload["excluded"] if i["code_id"])
    assert missing["issues"] == ["Quotation not found once in the transcript"]
    # The unlocated "1. Need" follows "B. Deliverers", so it is the deliverers' need.
    codes = {c["id"]: c for c in payload["codebook"]["codes"]}
    assert codes[missing["code_id"]]["parent_label"] == "B. Deliverers"
    assert any("Summary" in note for note in payload["notices"])
    assert payload["independent_human_coding"] is False

    path = save_reference_coding(project, payload, {transcript: workbook})
    assert list_references(project) == [path]
    saved = load_reference_coding(path, root)
    assert (path.parent / saved["files"][transcript]["copy"]).exists()
    with pytest.raises(ValueError, match="already exists"):
        save_reference_coding(project, payload, {transcript: workbook})
    (root / transcript).write_text("Changed.", encoding="utf-8")
    with pytest.raises(ValueError, match="transcript changed"):
        load_reference_coding(path, root)
    assert reference_coding_set(saved)["set_kind"] == "reference"


def test_reference_import_needs_an_id_name_and_statement(tmp_path):
    root, project, transcript = _cfir_study(tmp_path)
    for reference_id, label, statement in [("Bad ID", "x", "y"), ("ok-id", "", "y"), ("ok-id", "x", " ")]:
        with pytest.raises(ValueError):
            build_reference_coding(repo_root=root, project_path=project, reference_id=reference_id, label=label,
                                   statement=statement, files={}, created_by="R")


def test_command_line_help_runs():
    result = subprocess.run([sys.executable, str(ROOT / "scripts/qualbench.py"), "--help"],
                            capture_output=True, text=True, check=False)
    assert result.returncode == 0 and "compare" in result.stdout


def test_qualbench_page_builds_and_shows_a_report(tmp_path):
    from qualframe.dashboard import DashboardApp
    from qualframe.evidence_review import write_evidence_review_index
    from qualframe.pipeline.project_runner import ProjectRunner
    project, run_id, _, _ = _completed_synthetic_run(tmp_path)
    runner = ProjectRunner(project)
    plan = runner.build_run_plan()
    runner.save_run_plan(plan, "second-run")
    runner.execute_run_plan(plan, "second-run")
    for run in (run_id, "second-run"):
        write_evidence_review_index(project_manifest_path=project, run_id=run)
    app = DashboardApp(ProjectRunner(project))
    from qualframe.qualbench_ui import render_qualbench
    page = render_qualbench(app)
    assert "Build report" in page and 'name="run"' in page and "No earlier coding" in page
    with pytest.raises(ValueError, match="at least two"):
        app.handle_action("/qualbench/build", {"study_identity": [app.runner.project.project_id], "run": [run_id]})
    stamp, message = app.handle_action("/qualbench/build", {"study_identity": [app.runner.project.project_id],
                                                             "run": [run_id, "second-run"]})
    assert "QualBench report" in message
    app_page = render_qualbench(app, stamp)
    assert "Which codes were applied" in app_page and "Stability across repeats" in app_page
    assert "/file?" in app_page and "presence.csv" in app_page
    with pytest.raises(ValueError):
        render_qualbench(app, "../../etc")


def test_workbooks_are_paired_with_transcripts_by_id(tmp_path):
    from qualframe.qualbench_ui import match_workbooks
    transcripts = [tmp_path / "HE/12A_Transcript[cleaned]_CIHRAMN.docx", tmp_path / "HE/13A_Transcript.docx",
                   tmp_path / "ON/1A_Transcript.docx", tmp_path / "Interview with Ana.txt"]
    paired = match_workbooks(transcripts, ["Jayne 12A.xlsx", "CFIR - 13A - Zack.xlsx", "1A_CFIR_AllDomains.xlsx",
                                           "interview-with-ana coding.xlsx"])
    assert {k: v.name for k, v in paired.items()} == {
        "Jayne 12A.xlsx": "12A_Transcript[cleaned]_CIHRAMN.docx", "CFIR - 13A - Zack.xlsx": "13A_Transcript.docx",
        "1A_CFIR_AllDomains.xlsx": "1A_Transcript.docx", "interview-with-ana coding.xlsx": "Interview with Ana.txt"}
    with pytest.raises(ValueError, match="no transcript"):
        match_workbooks(transcripts, ["Summary.xlsx"])
    with pytest.raises(ValueError, match="more than one workbook"):
        match_workbooks(transcripts, ["Jayne 12A.xlsx", "Michelle 12A.xlsx"])


def test_reference_coding_can_be_imported_from_the_qualbench_page(tmp_path):
    from qualframe.dashboard import DashboardApp
    from qualframe.pipeline.project_runner import ProjectRunner
    from qualframe.qualbench_ui import render_qualbench
    from qualframe.web_forms import UploadedFile
    root, project, transcript = _cfir_study(tmp_path)
    app = DashboardApp(ProjectRunner(project))
    assert "Import an earlier coding" in render_qualbench(app)
    coding = Workbook(); sheet = coding.active; sheet.title = "Inner Setting"
    sheet.append(("CFIR Construct", "Quote", "Explanation"))
    sheet.append(("A. Structure", "We met every week in the old library.", "Neutral"))
    path = tmp_path / "upload.xlsx"; coding.save(path)
    upload = UploadedFile(filename="1A earlier coding.xlsx", content_type="application/octet-stream", content=path.read_bytes())
    fields = {"label": ["Earlier coding"], "statement": ["Model output reviewed by the team."]}
    message = app.handle_upload("/import/reference-coding", fields, {"upload": [upload]})
    assert "Imported Earlier coding: 1 transcript, 1 quotations found" in message
    saved = list_references(project)
    assert [p.parent.name for p in saved] == ["earlier-coding"]
    assert not any((project.parent / "references" / ".staging").iterdir())
    page = render_qualbench(app)
    assert "Earlier coding" in page and "Model output reviewed by the team." in page
    app.handle_upload("/import/reference-coding", fields, {"upload": [upload]})
    assert [p.parent.name for p in list_references(project)] == ["earlier-coding", "earlier-coding-2"]
    with pytest.raises(ValueError, match="name"):
        app.handle_upload("/import/reference-coding", {"label": [""], "statement": ["x"]}, {"upload": [upload]})


def test_runs_from_another_study_are_matched_by_transcript_content():
    here = _set("here", [("t/one.txt", "p", 0, 10)])
    there = _set("there", [("other/one copy.txt", "p", 0, 10)])
    there["sources"] = {"other/one copy.txt": "h1", "other/two copy.txt": "h2"}
    report = build_qualbench([here, there])
    assert report["transcripts"] == ["t/one.txt", "t/two.txt"]
    overall = next(r for r in report["presence"] if r["scope_kind"] == "all" and r["level"] == "code")
    assert overall["both"] == 1 and overall["percent_agreement"] == 1.0
    passages = next(r for r in report["passages"] if r["scope_kind"] == "all")
    assert passages["matched_exact"] == 1
