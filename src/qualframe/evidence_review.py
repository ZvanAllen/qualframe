from __future__ import annotations

import json
import html
import re
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from qualframe.config import load_codebook_spec, load_project_manifest, resolve_repo_path
from qualframe.directed_analysis import validate_evidence_span_against_source
from qualframe.evidence_anchoring import anchor_exact_quote
from qualframe.cfir_review import cfir_records, reference_codes
from qualframe.governance import assert_export_allowed
from qualframe.io.files import sha256_for_file, sha256_for_text
from qualframe.io.transcript_io import transcript_to_text
from qualframe.models import (
    CodeAssignment,
    EvidenceReviewAction,
    EvidenceReviewRecord,
    EvidenceSpan,
    MeaningUnit,
    RunManifest,
    SourceImportRecord,
    StageArtifact,
)


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _review_project(project_path: Path, artifact_root: Path):
    """Resolve the codebook retained by the run lock, not a later active version."""
    project = load_project_manifest(project_path)
    lock_path = artifact_root / 'analysis_configuration_lock.json'
    if lock_path.exists():
        lock = json.loads(lock_path.read_text(encoding='utf-8'))
        for asset in lock.get('assets', []):
            if asset.get('kind') == 'codebook' and asset.get('path'):
                path = resolve_repo_path(project_path.parents[2], asset['path'])
                if sha256_for_file(path) != asset.get('sha256'):
                    raise ValueError('The run codebook changed; restore its retained version before preparing review.')
                if path.suffix == '.xlsx':
                    project.paths['reference'] = str(path.parent)
                else:
                    project.codebook = str(path)
    return project


def _context_offsets(source_text: str, start: int, end: int, before: int, after: int, unit: str = "sentences") -> tuple[int, int]:
    if before == 0 and after == 0:
        return start, end
    if unit == "characters":
        return max(0, start - before), min(len(source_text), end + after)
    pattern = r"[^\n]+" if unit == "paragraphs" else r"[^.!?\n]+(?:[.!?]+|(?=\n|$))"
    ranges = [
        (match.start(), match.end())
        for match in re.finditer(pattern, source_text)
        if match.group(0).strip()
    ]
    overlapping = [i for i, (left, right) in enumerate(ranges) if left < end and right > start]
    if not overlapping:
        return start, end
    left = min(start, ranges[max(0, overlapping[0] - before)][0]) if before else start
    right = max(end, ranges[min(len(ranges) - 1, overlapping[-1] + after)][1]) if after else end
    while left < start and source_text[left].isspace():
        left += 1
    return left, right


def _source_records(repo_root: Path, project_path: Path) -> list[SourceImportRecord]:
    project = load_project_manifest(project_path)
    if not project.source_manifest:
        return []
    path = resolve_repo_path(repo_root, project.source_manifest)
    if not path.exists():
        return []
    return [
        SourceImportRecord.model_validate_json(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def _source_identity(
    repo_root: Path,
    source_path: Path,
    records: list[SourceImportRecord],
) -> tuple[str, str]:
    resolved = source_path.resolve()
    for record in records:
        destination = resolve_repo_path(repo_root, record.destination_path).resolve()
        if destination == resolved:
            return record.source_id, record.sha256
    digest = sha256_for_file(resolved)
    return f"source-{digest[:24]}", digest


def build_evidence_review_records(
    *,
    project_manifest_path: Path,
    run_id: str,
) -> list[EvidenceReviewRecord]:
    project_path = project_manifest_path.resolve()
    repo_root = project_path.parents[2]
    project = load_project_manifest(project_path)
    artifact_root = resolve_repo_path(repo_root, project.paths["artifacts"]) / run_id
    project = _review_project(project_path, artifact_root)
    manifest_path = artifact_root / "run_manifest.json"
    manifest = RunManifest.model_validate_json(manifest_path.read_text(encoding="utf-8"))
    codebook = (
        load_codebook_spec(resolve_repo_path(repo_root, project.codebook))
        if project.codebook
        else None
    )
    code_lookup = {entry.id: entry for entry in codebook.codes} if codebook else {}
    method_before = 0
    method_after = 0
    method_unit = "sentences"
    if project.method:
        import yaml

        method_data = yaml.safe_load(
            resolve_repo_path(repo_root, project.method).read_text(encoding="utf-8")
        )
        method_before = int(method_data.get("context", {}).get("before", 0))
        method_unit = method_data.get("context", {}).get("unit", "sentences")
        method_after = int(method_data.get("context", {}).get("after", 0))
    records = _source_records(repo_root, project_path)
    review_records: list[EvidenceReviewRecord] = []
    for artifact_rel in manifest.artifact_paths:
        artifact_path = artifact_root / artifact_rel
        artifact = StageArtifact.model_validate_json(artifact_path.read_text(encoding="utf-8"))
        if not artifact.parsed_output or not artifact.input_path:
            continue
        source_path = resolve_repo_path(repo_root, artifact.input_path)
        if not source_path.exists() or source_path.suffix.lower() not in {".txt", ".docx"}:
            continue
        source_text = transcript_to_text(source_path)
        source_text_hash = sha256_for_text(source_text)
        source_id, registered_hash = _source_identity(repo_root, source_path, records)
        current_hash = sha256_for_file(source_path)
        if artifact.stage_id == 'transcript_to_cfir':
            reference = resolve_repo_path(repo_root, project.paths['reference']) / 'codebook.xlsx'
            review_records.extend(cfir_records(artifact=artifact, artifact_rel=artifact_rel,
                artifact_path=artifact_path, source_path=source_path, repo_root=repo_root, manifest=manifest,
                codes=reference_codes(reference), codebook_hash=sha256_for_file(reference) if reference.exists() else 'unknown'))
            continue
        for sheet_name, rows in artifact.parsed_output.items():
            for row_index, row in enumerate(rows, start=1):
                if not {"Evidence Start", "Evidence End", "Exact Evidence", "Code ID"} <= set(row):
                    continue
                issues: list[str] = []
                raw_start = str(row.get("Evidence Start", "")).strip()
                raw_end = str(row.get("Evidence End", "")).strip()
                exact_text = str(row.get("Exact Evidence", ""))
                if artifact.evidence_location == "unique_exact_quote":
                    exact_text = html.unescape(exact_text)
                anchoring_method = "model_offsets"
                anchor_match_count = None
                automatic = artifact.evidence_location == "unique_exact_quote" and raw_start.lower() == raw_end.lower() == "auto"
                if automatic:
                    anchor = anchor_exact_quote(source_text, exact_text)
                    anchor_match_count = anchor.match_count
                    start, end = anchor.start if anchor.start is not None else -1, anchor.end if anchor.end is not None else -1
                    anchoring_method = "unique_exact_quote" if anchor.match_count == 1 else "unresolved_exact_quote"
                    if anchor.match_count != 1:
                        issues.append("Exact quotation is missing or occurs more than once; locate evidence through review.")
                else:
                    try:
                        start, end = int(raw_start), int(raw_end)
                    except ValueError:
                        start, end = -1, -1
                        issues.append("Evidence offsets are not integers under this run's location contract.")
                proposed_start, proposed_end = (-1, -1) if automatic else (start, end)
                meaning_unit_id = str(row.get("Meaning Unit ID", "")).strip() or (
                    f"meaning-{sha256_for_text(f'{artifact_rel}:{sheet_name}:{row_index}')[:16]}"
                )
                code_id = str(row.get("Code ID", "")).strip()
                code = code_lookup.get(code_id)
                if registered_hash != current_hash:
                    issues.append("Current source hash differs from the registered import hash.")
                if code is None:
                    issues.append(f"Code '{code_id}' is not active in the selected codebook.")
                span: EvidenceSpan | None = None
                assignment: CodeAssignment | None = None
                context_text: str | None = None
                if start < 0 or end <= start:
                    issues.append("Evidence offsets must specify a nonempty character span with 0 <= start < end.")
                if not exact_text.strip():
                    issues.append("Exact evidence text is empty.")
                if start >= 0 and end > start and exact_text.strip():
                    context_start, context_end = _context_offsets(
                        source_text, start, end, method_before, method_after, method_unit
                    )
                    try:
                        MeaningUnit(
                            id=meaning_unit_id,
                            source_id=source_id,
                            start_offset=start,
                            end_offset=end,
                            segmentation_method="model_proposed",
                        )
                        span = EvidenceSpan(
                            id=f"span-{sha256_for_text(f'{artifact_rel}:{sheet_name}:{row_index}')[:20]}",
                            meaning_unit_id=meaning_unit_id,
                            source_id=source_id,
                            source_sha256=source_text_hash,
                            start_offset=start,
                            end_offset=end,
                            exact_text=exact_text,
                            context_start_offset=context_start,
                            context_end_offset=context_end,
                            creation_method="automatic" if automatic else "model_proposed",
                        )
                        validate_evidence_span_against_source(span, source_text)
                    except ValueError as exc:
                        issues.append(str(exc))
                        span = None
                    else:
                        context_text = source_text[context_start:context_end]
                        if code is not None and codebook is not None:
                            assignment = CodeAssignment(
                                id=f"assignment-{sha256_for_text(f'{artifact_rel}:{sheet_name}:{row_index}:{code_id}')[:20]}",
                                meaning_unit_id=meaning_unit_id,
                                evidence_span_id=span.id,
                                codebook_id=codebook.id,
                                codebook_version=codebook.version,
                                code_id=code_id,
                                origin="model",
                                rationale=str(row.get("Rationale", "")).strip() or None,
                                run_id=run_id,
                            )
                record_id = f"evidence-{sha256_for_text(f'{artifact_rel}:{sheet_name}:{row_index}')[:20]}"
                review_records.append(
                    EvidenceReviewRecord(
                        record_id=record_id,
                        run_id=run_id,
                        stage_id=artifact.stage_id,
                        artifact_path=artifact_rel,
                        artifact_sha256=sha256_for_file(artifact_path),
                        sheet_name=sheet_name,
                        row_index=row_index,
                        source_id=source_id,
                        source_path=str(source_path.relative_to(repo_root)),
                        source_sha256=current_hash,
                        meaning_unit_id=meaning_unit_id,
                        proposed_start_offset=proposed_start,
                        proposed_end_offset=proposed_end,
                        raw_start_value=raw_start, raw_end_value=raw_end,
                        anchoring_method=anchoring_method, anchor_match_count=anchor_match_count,
                        proposed_exact_text=exact_text,
                        context_text=context_text,
                        evidence_span=span,
                        assignment=assignment,
                        proposed_code_id=code_id, proposed_rationale=str(row.get("Rationale", "")),
                        code_label=code.label if code else None,
                        code_definition=code.definition if code else None,
                        model=artifact.model,
                        provider=artifact.provider,
                        prompt_sha256=artifact.prompt_sha256,
                        configuration_fingerprint=manifest.configuration_fingerprint,
                        validation_issues=sorted(set(issues)),
                    )
                )
    return review_records


def write_evidence_review_index(
    *, project_manifest_path: Path, run_id: str
) -> Path:
    project_path = project_manifest_path.resolve()
    repo_root = project_path.parents[2]
    project = load_project_manifest(project_path)
    artifact_root = resolve_repo_path(repo_root, project.paths["artifacts"]) / run_id
    project = _review_project(project_path, artifact_root)
    output = artifact_root / "evidence_review" / "index.json"
    if output.exists():
        for record in load_evidence_review_records(output):
            if sha256_for_file(artifact_root / record.artifact_path) != record.artifact_sha256:
                raise ValueError('Model artifact changed after review preparation; preserve this index and create a new run.')
        return output  # Review indexes are immutable once materialized.
    records = build_evidence_review_records(
        project_manifest_path=project_path, run_id=run_id
    )
    output = artifact_root / "evidence_review" / "index.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(
            {
                "run_id": run_id,
                "generated_at_utc": _utc_now(),
                "record_count": len(records),
                "codebook": (load_codebook_spec(resolve_repo_path(repo_root, project.codebook)).model_dump(mode='json') if project.codebook else
                    {'id': 'cfir-reference', 'version': sha256_for_file(resolve_repo_path(repo_root, project.paths['reference']) / 'codebook.xlsx') if (resolve_repo_path(repo_root, project.paths['reference']) / 'codebook.xlsx').exists() else 'unknown',
                     'codes': reference_codes(resolve_repo_path(repo_root, project.paths['reference']) / 'codebook.xlsx')}),
                "records": [record.model_dump(mode="json") for record in records],
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    return output


def load_evidence_review_records(index_path: Path) -> list[EvidenceReviewRecord]:
    payload = json.loads(index_path.read_text(encoding="utf-8"))
    return [EvidenceReviewRecord.model_validate(item) for item in payload.get("records", [])]


def load_evidence_review_actions(review_dir: Path) -> list[EvidenceReviewAction]:
    path = review_dir / "evidence" / "actions.jsonl"
    if not path.exists():
        return []
    return [
        EvidenceReviewAction.model_validate_json(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def latest_evidence_action_by_record(review_dir: Path) -> dict[str, EvidenceReviewAction]:
    latest: dict[str, EvidenceReviewAction] = {}
    for action in load_evidence_review_actions(review_dir):
        latest[action.record_id] = action
    return latest


def blocking_review_issues(record: EvidenceReviewRecord, source_text: str) -> list[str]:
    """Accept itself confirms uniquely located typography-only differences."""
    issues = list(record.validation_issues)
    legacy_notice = 'Formatting differs from the model quotation. Confirm the highlighted source passage with Save changes.'
    if record.anchoring_method == 'normalized_quote_candidate' and record.evidence_span:
        from qualframe.cfir_review import locate_quote
        span = record.evidence_span
        start, end, method = locate_quote(source_text, record.proposed_exact_text)
        if (start, end) == (span.start_offset, span.end_offset) and method == 'normalized_quote_candidate':
            issues = [issue for issue in issues if issue != legacy_notice]
    return issues


def append_evidence_review_action(
    *,
    record: EvidenceReviewRecord,
    source_path: Path,
    review_dir: Path,
    artifact_root: Path,
    decision: str,
    reviewer: str,
    rationale: str,
    edited_start: int | None = None,
    edited_end: int | None = None,
    edited_code_id: str | None = None,
    codebook_id: str | None = None,
    codebook_version: str | None = None,
) -> EvidenceReviewAction:
    if decision not in {"accept", "edit", "reject", "defer"}:
        raise ValueError("Unknown evidence review decision.")
    if not reviewer.strip() or not rationale.strip():
        raise ValueError("Reviewer and rationale are required.")
    if sha256_for_file(source_path) != record.source_sha256:
        raise ValueError("Source file hash changed after the evidence review index was created.")
    source_text = transcript_to_text(source_path)
    issues = blocking_review_issues(record, source_text)
    if decision == "accept" and record.evidence_span:
        validate_evidence_span_against_source(record.evidence_span, source_text)
    if decision == "accept" and record.record_kind != 'absence' and (
        record.evidence_span is None or record.assignment is None or issues
    ):
        raise ValueError("Invalid evidence cannot be accepted; edit, reject, or defer it.")
    if decision == 'accept' and record.record_kind == 'absence' and issues:
        raise ValueError('Resolve the unknown construct before accepting an absence proposal.')
    action_id = f"evidence-review-{uuid4().hex}"
    derived_rel: str | None = None
    if decision == "edit":
        if edited_start is None or edited_end is None or not edited_code_id:
            raise ValueError("Edit requires start, end, and code ID.")
        if edited_start < 0 or edited_end <= edited_start or edited_end > len(source_text):
            raise ValueError("Edited offsets are outside the immutable source text.")
        exact_text = source_text[edited_start:edited_end]
        span = EvidenceSpan(
            id=f"span-{uuid4().hex}",
            meaning_unit_id=record.meaning_unit_id,
            source_id=record.source_id,
            source_sha256=sha256_for_text(source_text),
            start_offset=edited_start,
            end_offset=edited_end,
            exact_text=exact_text,
            context_start_offset=edited_start,
            context_end_offset=edited_end,
            creation_method="reviewer_edited",
            derived_from_span_id=(record.evidence_span.id if record.evidence_span else None),
        )
        validate_evidence_span_against_source(span, source_text)
        original_assignment = record.assignment
        assignment = CodeAssignment(
            id=f"assignment-{uuid4().hex}",
            meaning_unit_id=record.meaning_unit_id,
            evidence_span_id=span.id,
            codebook_id=codebook_id or (original_assignment.codebook_id if original_assignment else "reviewer_selected"),
            codebook_version=codebook_version or (original_assignment.codebook_version if original_assignment else "unknown"),
            code_id=edited_code_id,
            origin="adjudicated",
            rationale=rationale.strip(),
            run_id=record.run_id,
            review_status="edited",
            derived_from_assignment_id=(original_assignment.id if original_assignment else None),
        )
        derived_path = artifact_root / "reviewed_evidence" / f"{action_id}.json"
        derived_path.parent.mkdir(parents=True, exist_ok=True)
        derived_path.write_text(
            json.dumps(
                {
                    "kind": "reviewed_evidence_edit",
                    "record_id": record.record_id,
                    "record_sha256": sha256_for_text(record.model_dump_json()),
                    "evidence_span": span.model_dump(mode="json"),
                    "assignment": assignment.model_dump(mode="json"),
                    "created_at_utc": _utc_now(),
                },
                indent=2,
            ),
            encoding="utf-8",
        )
        derived_rel = str(derived_path.relative_to(artifact_root))
    action = EvidenceReviewAction(
        action_id=action_id,
        run_id=record.run_id,
        record_id=record.record_id,
        record_sha256=sha256_for_text(record.model_dump_json()),
        decision=decision,
        reviewer=reviewer.strip(),
        rationale=rationale.strip(),
        derived_artifact_path=derived_rel,
        timestamp_utc=_utc_now(),
    )
    log_path = review_dir / "evidence" / "actions.jsonl"
    log_path.parent.mkdir(parents=True, exist_ok=True)
    with log_path.open("a", encoding="utf-8") as handle:
        handle.write(action.model_dump_json() + "\n")
    return action


def export_evidence_review_package(
    *,
    project_manifest_path: Path,
    run_id: str,
    require_released: bool = False,
) -> list[Path]:
    project_path = project_manifest_path.resolve()
    repo_root = project_path.parents[2]
    project = load_project_manifest(project_path)
    artifact_root = resolve_repo_path(repo_root, project.paths["artifacts"]) / run_id
    assert_export_allowed(artifact_root, require_released)
    index_path = artifact_root / "evidence_review" / "index.json"
    if not index_path.exists():
        raise ValueError("Prepare the exact-evidence review index before exporting.")
    records = load_evidence_review_records(index_path)
    review_dir = resolve_repo_path(repo_root, project.paths["review_queue"]) / run_id
    actions = load_evidence_review_actions(review_dir)
    latest = {action.record_id: action for action in actions}
    unresolved = [
        record.record_id
        for record in records
        if record.record_id not in latest or latest[record.record_id].decision == "defer"
    ]
    if unresolved:
        raise PermissionError(
            f"Reviewed evidence export requires a resolved action for every item; "
            f"{len(unresolved)} item(s) remain unreviewed or deferred."
        )
    manifest = json.loads((artifact_root / "run_manifest.json").read_text(encoding="utf-8"))
    catalog = json.loads(index_path.read_text(encoding="utf-8")).get("codebook", {})
    labels = {c["id"]: c["label"] for c in catalog.get("codes", [])}
    evidence_rows: list[dict[str, object]] = []
    for record in records:
        action = latest.get(record.record_id)
        span = record.evidence_span
        assignment = record.assignment
        if action and action.derived_artifact_path:
            derived = json.loads(
                (artifact_root / action.derived_artifact_path).read_text(encoding="utf-8")
            )
            span = EvidenceSpan.model_validate(derived["evidence_span"])
            assignment = CodeAssignment.model_validate(derived["assignment"])
        evidence_rows.append(
            {
                "Record ID": record.record_id,
                "Review decision": action.decision if action else "unreviewed",
                "Source ID": record.source_id,
                "Source path": record.source_path,
                "Source SHA-256": record.source_sha256,
                "Source text SHA-256": span.source_sha256 if span else "",
                "Start offset": span.start_offset if span else record.proposed_start_offset,
                "End offset": span.end_offset if span else record.proposed_end_offset,
                "Exact evidence": span.exact_text if span else record.proposed_exact_text,
                "Context": record.context_text or "",
                "Code ID": assignment.code_id if assignment else record.proposed_code_id or "",
                "Code label": labels.get(assignment.code_id if assignment else record.proposed_code_id, record.code_label or ""),
                "Proposal kind": record.record_kind,
                "Original model quotation": record.proposed_exact_text,
                "Rationale": assignment.rationale if assignment else "",
                "Validation issues": "; ".join(record.validation_issues),
                "Stage": record.stage_id,
                "Anchoring method": record.anchoring_method,
                "Raw model start": record.raw_start_value or str(record.proposed_start_offset),
                "Raw model end": record.raw_end_value or str(record.proposed_end_offset),
                "Artifact": record.artifact_path,
                "Artifact SHA-256": record.artifact_sha256,
                "Prompt SHA-256": record.prompt_sha256,
                "Configuration fingerprint": record.configuration_fingerprint or "",
            }
        )
    package = {
        "kind": "qualframe_reviewed_evidence_export",
        "status": "researcher_review_record_not_automatic_findings",
        "generated_at_utc": _utc_now(),
        "project_id": project.project_id,
        "data_classification": project.data_classification,
        "run_manifest": manifest,
        "evidence": evidence_rows,
        "review_actions": [action.model_dump(mode="json") for action in actions],
    }
    output_dir = resolve_repo_path(repo_root, project.paths["exports"]) / run_id
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "reviewed_evidence_package.json"
    json_path.write_text(json.dumps(package, indent=2), encoding="utf-8")

    workbook_path = output_dir / "reviewed_evidence_package.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "Evidence", evidence_rows)
    _write_sheet(
        workbook,
        "Review History",
        [action.model_dump(mode="json") for action in actions],
    )
    methods_rows = [
        {"Field": "Project ID", "Value": project.project_id},
        {"Field": "Run ID", "Value": run_id},
        {"Field": "Framework", "Value": manifest.get("framework_id", "")},
        {"Field": "Pipeline", "Value": manifest.get("pipeline_id", "")},
        {"Field": "Method", "Value": manifest.get("method_id", "")},
        {"Field": "Codebook", "Value": manifest.get("codebook_id", "")},
        {"Field": "Configuration fingerprint", "Value": manifest.get("configuration_fingerprint", "")},
        {"Field": "Provider", "Value": project.provider.name},
        {"Field": "Model", "Value": project.provider.model},
        {"Field": "Data classification", "Value": project.data_classification or ""},
        {"Field": "Interpretation boundary", "Value": "Reviewed coding evidence; not automatic findings."},
    ]
    _write_sheet(workbook, "Methods", methods_rows)
    workbook.save(workbook_path)
    return [json_path, workbook_path]


def _write_sheet(workbook: Workbook, title: str, rows: list[dict[str, object]]) -> None:
    worksheet = workbook.create_sheet(title=title[:31])
    if not rows:
        worksheet.append(["No records"])
        return
    headers = list(rows[0])
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAD3")
    for row in rows:
        worksheet.append(
            [
                json.dumps(row.get(header), ensure_ascii=False)
                if isinstance(row.get(header), (dict, list))
                else row.get(header, "")
                for header in headers
            ]
        )
    for column in worksheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
        worksheet.column_dimensions[column[0].column_letter].width = max(width, 12)
        for cell in column:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    worksheet.freeze_panes = "A2"
