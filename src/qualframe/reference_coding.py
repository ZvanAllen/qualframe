"""Import an earlier CFIR coding of the study's transcripts as a reference for comparison.

The earlier coding comes as workbooks in the published CFIR table format: one sheet per domain
with the columns CFIR Construct, Quote and Explanation (CFIR Domain, Page Number and reviewer
note columns are optional). Each construct label is matched to the study's CFIR codebook and each
quotation is located in the study's copy of the transcript. Nothing is changed in the original
workbooks; copies are kept with the reference so it can be checked later.

A reference is not treated as correct coding. Its provenance statement says what it is (for
example, model output reviewed by a team) and is carried into every comparison that uses it.
"""
from __future__ import annotations

import json
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path

from qualframe.cfir_review import (
    REVIEW_LABEL_POLICY,
    _outline_candidates,
    locate_quote,
    normalized_with_offsets,
    reference_codes,
)
from qualframe.config import load_project_manifest, resolve_repo_path
from qualframe.io.files import sha256_for_file, sha256_for_text
from qualframe.io.transcript_io import transcript_to_text

KIND = "qualframe_reference_coding_v1"
NO_QUOTE = "no relevant quote found"
_DOMAINS = (
    ("outcome", "Outcomes Addendum"),
    ("outer", "Outer Setting"),
    ("inner", "Inner Setting"),
    ("individual", "Individuals"),
    ("process", "Implementation Process"),
    ("innovation", "Innovation"),
)
_STANDARD_COLUMNS = {"cfir construct", "quote", "explanation", "page number", "page", "cfir domain"}


def sheet_domain(title: str) -> str | None:
    """CFIR domain named by a sheet title, such as "Inner Setting" or "CFIR Outcomes"."""
    lowered = title.casefold()
    for key, domain in _DOMAINS:
        if key in lowered:
            return domain
    return None


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_cfir_coding_workbook(path: Path) -> tuple[list[dict], list[str]]:
    """Coding rows from every domain sheet, and notes about sheets that were skipped."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows: list[dict] = []
    skipped: list[str] = []
    try:
        for sheet in workbook:
            domain = sheet_domain(sheet.title)
            if domain is None:
                skipped.append(f"Sheet '{sheet.title}' is not a CFIR domain.")
                continue
            header: dict[str, int] | None = None
            found = False
            for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                cells = [_cell(value) for value in values]
                lowered = [cell.casefold() for cell in cells]
                if "cfir construct" in lowered and "quote" in lowered:
                    header = {name: index for index, name in enumerate(lowered) if name}
                    found = True
                    continue
                if header is None:
                    continue
                construct = cells[header["cfir construct"]] if header["cfir construct"] < len(cells) else ""
                quote = cells[header["quote"]] if header["quote"] < len(cells) else ""
                if not construct:
                    continue
                notes = {
                    name: cells[index]
                    for name, index in header.items()
                    if name not in _STANDARD_COLUMNS and index < len(cells) and cells[index]
                }
                explanation_index = header.get("explanation")
                rows.append(dict(
                    sheet=sheet.title, row=number, domain=domain, construct=construct, quote=quote,
                    explanation=cells[explanation_index] if explanation_index is not None and explanation_index < len(cells) else "",
                    notes=notes,
                ))
            if not found:
                skipped.append(f"Sheet '{sheet.title}' has no CFIR Construct and Quote columns.")
    finally:
        workbook.close()
    return rows, skipped


def locate_reference_quote(source: str, quote: str) -> tuple[int, int, str]:
    """Locate a quotation; quotations shortened with an ellipsis are located by their parts."""
    start, end, method = locate_quote(source, quote)
    if start >= 0:
        return start, end, method
    parts = [part.strip(" \"'“”‘’") for part in re.split(r"\.\.\.|…", quote)]
    parts = [part for part in parts if len(part) >= 12]
    if len(parts) < 2:
        return -1, -1, "unresolved_exact_quote"
    normalized, offsets = normalized_with_offsets(source)
    position, spans = 0, []
    for part in parts:
        needle = normalized_with_offsets(part)[0]
        found = normalized.find(needle, position)
        if found < 0 or normalized.find(needle, found + 1) >= 0:
            return -1, -1, "unresolved_exact_quote"
        spans.append((offsets[found][0], offsets[found + len(needle) - 1][1]))
        position = found + len(needle)
    return spans[0][0], spans[-1][1], "ellipsis_parts"


def _reference_quote_anchor(source: str, quote: str) -> dict | None:
    """Keep existing unique anchors, otherwise retain repeated whole-quote alternatives.

    Alternatives describe one annotation with an unresolved location, not separate coded
    passages. Exact occurrences take precedence over normalized ones, as for unique matches.
    The ellipsis locator above is unchanged; no alternative combinations of its parts are made.
    """
    start, end, method = locate_reference_quote(source, quote)
    if start >= 0:
        return dict(start=start, end=end, exact_text=source[start:end], anchoring_method=method)
    needle = quote.strip()
    if len(needle) > 1 and (needle[0], needle[-1]) in [('"', '"'), ('“', '”')]:
        needle = needle[1:-1]
    if not needle.strip():
        return None

    def occurrences(text: str, fragment: str) -> list[tuple[int, int]]:
        if not fragment:
            return []
        found = []
        position = text.find(fragment)
        while position >= 0:
            found.append((position, position + len(fragment)))
            position = text.find(fragment, position + 1)
        return found

    spans = occurrences(source, needle)
    method = "exact_repeated"
    if len(spans) < 2:
        normalized, offsets = normalized_with_offsets(source)
        spans = [(offsets[a][0], offsets[b - 1][1])
                 for a, b in occurrences(normalized, normalized_with_offsets(needle)[0])]
        method = "normalized_repeated"
    spans = list(dict.fromkeys(spans))
    if len(spans) < 2:
        return None
    return dict(match="repeated", anchoring_method=method,
                alternative_spans=[dict(start=a, end=b, exact_text=source[a:b]) for a, b in spans])


def build_reference_coding(*, repo_root: Path, project_path: Path, reference_id: str, label: str,
                           statement: str, files: dict[str, Path], created_by: str,
                           independent_human_coding: bool = False) -> dict:
    """Build the reference from workbooks keyed by the transcript each one codes.

    ``files`` maps a transcript's workspace-relative path (as runs record it) to a workbook.
    """
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{2,63}", reference_id):
        raise ValueError("Use a short reference ID of lowercase letters, numbers and hyphens.")
    if not label.strip() or not statement.strip():
        raise ValueError("Give the reference a name and say what it is (how it was coded and reviewed).")
    repo_root = repo_root.resolve()
    project = load_project_manifest(project_path)
    workbook = resolve_repo_path(repo_root, project.paths["reference"]) / "codebook.xlsx"
    codes = reference_codes(workbook)
    if not codes:
        raise ValueError("Import the CFIR codebook workbook before importing a reference coding.")
    known = {code["id"] for code in codes}
    sources: dict[str, str] = {}
    provenance: dict[str, dict] = {}
    included, excluded, absences, notices = [], [], [], []
    for source_path, workbook_path in sorted(files.items()):
        source_key = source_path.replace("\\", "/")
        transcript = (repo_root / source_key).resolve()
        if not transcript.is_relative_to(repo_root) or not transcript.exists():
            raise ValueError(f"Transcript not found in this workspace: {source_key}")
        text = transcript_to_text(transcript)
        sources[source_key] = sha256_for_file(transcript)
        workbook_hash = sha256_for_file(workbook_path)
        provenance[source_key] = dict(workbook=Path(workbook_path).name, original_path=str(workbook_path),
                                      workbook_sha256=workbook_hash)
        rows, skipped = read_cfir_coding_workbook(Path(workbook_path))
        notices.extend(f"{Path(workbook_path).name}: {note}" for note in skipped)
        current_parent: dict[str, str] = {}
        for row in rows:
            in_domain = [c for c in codes if row["domain"].lower().replace(" addendum", "") in c["domain"].lower()]
            candidates = _outline_candidates(row["construct"], in_domain, current_parent.get(row["domain"], ""))
            code = candidates[0] if len(candidates) == 1 else None
            if code:
                current_parent[row["domain"]] = code["parent_label"] or code["raw_label"]
            identity = sha256_for_text(f"{workbook_hash}:{row['sheet']}:{row['row']}")[:20]
            item = dict(
                record_id=f"reference-{identity}", source=source_key, code_id=code["id"] if code else None,
                construct_label=row["construct"], domain=row["domain"], quote=row["quote"],
                explanation=row["explanation"], notes=row["notes"], sheet=row["sheet"], row=row["row"],
            )
            if not row["quote"]:
                continue  # a heading row in the table, such as a parent construct with no quotation
            if row["quote"].strip().strip("\"“”").casefold() == NO_QUOTE:
                if code:
                    absences.append(item)
                else:
                    excluded.append(dict(item, issues=["Construct not found in the codebook"]))
                continue
            anchor = _reference_quote_anchor(text, row["quote"])
            issues = []
            if code is None:
                issues.append("Construct not found in the codebook")
            if anchor is None:
                issues.append("Quotation not found once in the transcript")
            if issues:
                details = anchor if anchor and anchor.get("match") == "repeated" else dict(
                    anchoring_method=anchor["anchoring_method"] if anchor else "unresolved_exact_quote")
                excluded.append(dict(item, **details, issues=issues))
            else:
                included.append(dict(item, **anchor))
    payload = dict(
        kind=KIND,
        reference_id=reference_id,
        run_id=f"reference:{reference_id}",
        project_id=project.project_id,
        label=" ".join(label.split()),
        statement=" ".join(statement.split()),
        independent_human_coding=independent_human_coding,
        created_at_utc=datetime.now(timezone.utc).isoformat(),
        created_by=created_by,
        label_policy=REVIEW_LABEL_POLICY,
        codebook=dict(id="cfir-reference", version=sha256_for_file(workbook), codes=codes),
        sources=sources,
        files=provenance,
        included=included,
        excluded=excluded,
        absences=absences,
        notices=notices,
        counts=dict(included=len(included), excluded=len(excluded), absences=len(absences),
                    unmatched_constructs=sum(1 for item in excluded if item["code_id"] is None),
                    quotations_not_found=sum(1 for item in excluded if item["code_id"] in known)),
    )
    return payload


def reference_folder(project_path: Path, reference_id: str) -> Path:
    return project_path.resolve().parent / "references" / reference_id


def save_reference_coding(project_path: Path, payload: dict, files: dict[str, Path]) -> Path:
    """Write the reference and copies of its workbooks into the study. Never overwrites."""
    folder = reference_folder(project_path, payload["reference_id"])
    if folder.exists():
        raise ValueError("A reference with this ID already exists in the study. Choose another ID.")
    (folder / "workbooks").mkdir(parents=True)
    for source_key, workbook in files.items():
        entry = payload["files"][source_key.replace("\\", "/")]
        copy = folder / "workbooks" / f"{entry['workbook_sha256'][:12]}-{Path(workbook).name}"
        shutil.copy2(workbook, copy)
        entry["copy"] = copy.relative_to(folder).as_posix()
    path = folder / "reference.json"
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def list_references(project_path: Path) -> list[Path]:
    root = project_path.resolve().parent / "references"
    return sorted(root.glob("*/reference.json")) if root.exists() else []


def load_reference_coding(path: Path, repo_root: Path) -> dict:
    """Load a saved reference and check that its transcripts and workbook copies are unchanged."""
    payload = json.loads(path.read_text(encoding="utf-8"))
    if payload.get("kind") != KIND:
        raise ValueError("This file is not a QualFrame reference coding.")
    for source_key, digest in payload["sources"].items():
        if sha256_for_file(repo_root / source_key) != digest:
            raise ValueError(f"A transcript changed after the reference was imported: {source_key}")
    for entry in payload["files"].values():
        if entry.get("copy") and sha256_for_file(path.parent / entry["copy"]) != entry["workbook_sha256"]:
            raise ValueError(f"A reference workbook copy changed: {entry['copy']}")
    return payload


# ------------------------------------------------------------- coding tables

_DOCUMENT_COLUMNS = ("document", "transcript", "file", "filename", "source", "paper", "hbcp_doc_id")
_CODE_COLUMNS = ("code id", "code_id", "code", "code label", "construct", "cfir construct", "bct_id", "bct")
_QUOTE_COLUMNS = ("quotation", "quote", "passage", "evidence", "excerpt", "text", "annotated_text")


def read_coding_table(path: Path) -> tuple[list[dict], dict[str, str]]:
    """Rows of a coding table in CSV or Excel: a document column, a code column, an optional quotation column.

    Returns the rows and the columns used. Other columns are kept as notes.
    """
    import csv

    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            table = [[_cell(value) for value in row] for row in csv.reader(handle)]
    else:
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            table = [[_cell(value) for value in row] for row in next(iter(workbook)).iter_rows(values_only=True)]
        finally:
            workbook.close()
    table = [row for row in table if any(row)]
    if not table:
        raise ValueError("The coding table is empty.")
    header = [cell.casefold().strip() for cell in table[0]]

    def column(names: tuple[str, ...]) -> int | None:
        return next((header.index(name) for name in names if name in header), None)

    document, code, quote = column(_DOCUMENT_COLUMNS), column(_CODE_COLUMNS), column(_QUOTE_COLUMNS)
    if document is None or code is None:
        raise ValueError("The first row needs a Document column (the transcript's name) and a Code column; "
                         "a Quotation column is optional.")
    used = {"document": table[0][document], "code": table[0][code],
            **({"quotation": table[0][quote]} if quote is not None else {})}
    rows = []
    for number, row in enumerate(table[1:], 2):
        row = row + [""] * (len(header) - len(row))
        notes = {table[0][i]: row[i] for i in range(len(header))
                 if i not in {document, code, quote} and table[0][i] and row[i]}
        rows.append(dict(row=number, document=row[document], code=row[code],
                         quote=row[quote] if quote is not None else "", notes=notes))
    return rows, used


def _study_codes(repo_root: Path, project) -> list[dict]:
    from qualframe.config import load_codebook_spec

    if project.codebook:
        spec = load_codebook_spec(resolve_repo_path(repo_root, project.codebook))
        return [dict(id=code.id, label=code.label, raw_label=code.label, parent_id=code.parent_id,
                     active=code.active) for code in spec.codes]
    return reference_codes(resolve_repo_path(repo_root, project.paths["reference"]) / "codebook.xlsx")


def _transcript_lookup(repo_root: Path, project) -> list[Path]:
    from qualframe.io.transcript_io import discover_transcripts

    folder = resolve_repo_path(repo_root, project.paths.get("raw_transcripts", "data/raw/transcripts"))
    return [path for paths in discover_transcripts(folder).values() for path in paths]


def match_document(name: str, transcripts: list[Path]) -> Path | None:
    """The transcript a table names: by file name or name without extension, else by a leading ID such as 12A."""
    wanted = name.strip().casefold()
    exact = [t for t in transcripts if wanted in {t.name.casefold(), t.stem.casefold()}]
    if len(exact) == 1:
        return exact[0]
    tokens = {token.casefold() for token in re.findall(r"[A-Za-z0-9]+", name)}
    by_id = []
    for transcript in transcripts:
        first = re.match(r"[A-Za-z0-9]+", transcript.stem)
        if first and re.search(r"\d", first.group(0)) and first.group(0).casefold() in tokens:
            by_id.append(transcript)
    return by_id[0] if len(by_id) == 1 else None


def build_reference_from_table(*, repo_root: Path, project_path: Path, reference_id: str, label: str,
                               statement: str, table: Path, created_by: str,
                               independent_human_coding: bool = False) -> dict:
    """A reference coding from a table of document, code and (optionally) quotation rows.

    Codes are matched by ID, else by name (ignoring outline numbers such as "3.1"). Rows without a
    quotation count as the code being present in the document, with nothing to compare passages with.
    """
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{2,63}", reference_id):
        raise ValueError("Use a short reference ID of lowercase letters, numbers and hyphens.")
    if not label.strip() or not statement.strip():
        raise ValueError("Give the reference a name and say what it is (how it was coded and reviewed).")
    from qualframe.cfir_review import construct_name

    repo_root = repo_root.resolve()
    project = load_project_manifest(project_path)
    codes = _study_codes(repo_root, project)
    if not codes:
        raise ValueError("Add the study's codebook before importing an earlier coding.")
    by_id = {code["id"].casefold(): code for code in codes}
    by_name: dict[str, list[dict]] = {}
    for code in codes:
        for name in {construct_name(code["label"]), construct_name(code.get("raw_label", ""))} - {""}:
            by_name.setdefault(name, []).append(code)
    transcripts = _transcript_lookup(repo_root, project)
    rows, used = read_coding_table(table)
    texts: dict[Path, str] = {}
    sources: dict[str, str] = {}
    included, excluded, absences, presence_only = [], [], [], []
    unmatched_documents: set[str] = set()
    table_hash = sha256_for_file(table)
    for row in rows:
        transcript = match_document(row["document"], transcripts)
        candidates = [by_id[row["code"].casefold()]] if row["code"].casefold() in by_id else \
            list({c["id"]: c for c in by_name.get(construct_name(row["code"]), [])}.values())
        code = candidates[0] if len(candidates) == 1 else None
        identity = sha256_for_text(f"{table_hash}:{row['row']}")[:20]
        source_key = transcript.resolve().relative_to(repo_root).as_posix() if transcript else None
        item = dict(record_id=f"reference-{identity}", source=source_key, code_id=code["id"] if code else None,
                    construct_label=row["code"], document=row["document"], quote=row["quote"], notes=row["notes"],
                    row=row["row"])
        if transcript is None:
            unmatched_documents.add(row["document"])
            excluded.append(dict(item, issues=["Document not found in this study"]))
            continue
        if transcript not in texts:
            texts[transcript] = transcript_to_text(transcript)
            sources[source_key] = sha256_for_file(transcript)
        if code is None:
            excluded.append(dict(item, issues=["Code not found in the codebook" if not candidates
                                               else "Code name matches more than one code"]))
            continue
        quote = row["quote"].strip()
        if not quote:
            presence_only.append(item)
            continue
        if quote.strip("\"“”").casefold() == NO_QUOTE:
            absences.append(item)
            continue
        anchor = _reference_quote_anchor(texts[transcript], quote)
        if anchor is None:
            excluded.append(dict(item, anchoring_method="unresolved_exact_quote", issues=["Quotation not found once in the transcript"]))
        else:
            included.append(dict(item, **anchor))
    known = {code["id"] for code in codes}
    return dict(
        kind=KIND, reference_id=reference_id, run_id=f"reference:{reference_id}", project_id=project.project_id,
        label=" ".join(label.split()), statement=" ".join(statement.split()), independent_human_coding=independent_human_coding,
        created_at_utc=datetime.now(timezone.utc).isoformat(), created_by=created_by,
        label_policy=REVIEW_LABEL_POLICY, columns=used,
        codebook=dict(id=project.codebook or "cfir-reference", version=table_hash, codes=codes),
        sources=sources,
        files={"table": dict(workbook=table.name, original_path=str(table), workbook_sha256=table_hash)},
        included=included, excluded=excluded, absences=absences, presence_only=presence_only,
        notices=[f"Not in this study: {name}" for name in sorted(unmatched_documents)],
        counts=dict(included=len(included), excluded=len(excluded), absences=len(absences),
                    presence_only=len(presence_only),
                    unmatched_constructs=sum(1 for item in excluded if item.get("source") and item["code_id"] is None),
                    quotations_not_found=sum(1 for item in excluded if item.get("source") and item["code_id"] in known),
                    documents_not_found=len(unmatched_documents)),
    )
