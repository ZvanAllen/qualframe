"""Bounded, deterministic Excel codebook reading; no formula evaluation."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from qualframe.models import CodebookEntry, CodebookSpec

FIELDS = {
    "id": "Code ID",
    "label": "Code label",
    "definition": "Definition",
    "parent_id": "Parent code ID",
    "inclusion_guidance": "Include when",
    "exclusion_guidance": "Exclude when",
    "examples": "Examples (one per line)",
}
REQUIRED = {"id", "label", "definition"}


@dataclass(frozen=True)
class ExcelSheet:
    name: str
    rows: list[list[str]]


def read_excel(content: bytes) -> list[ExcelSheet]:
    if not content or len(content) > 25 * 1024 * 1024:
        raise ValueError("Choose an Excel workbook smaller than 25 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds the 100 MB import limit.")
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            result = []
            for sheet in workbook.worksheets:
                if sheet.max_row > 10001 or sheet.max_column > 100:
                    raise ValueError("Each sheet must have at most 10,001 rows and 100 columns.")
                rows = []
                for cells in sheet.iter_rows():
                    row = []
                    for cell in cells:
                        if cell.data_type in {"f", "e"}:
                            raise ValueError(
                                f"{sheet.title}!{cell.coordinate}: replace formulas/errors "
                                "with plain values before importing."
                            )
                        value = cell.value
                        if value is None:
                            row.append("")
                        elif isinstance(value, (str, int, float)) and not isinstance(value, bool):
                            row.append(str(value).strip())
                        else:
                            raise ValueError(
                                f"{sheet.title}!{cell.coordinate}: use text or numbers, "
                                "not dates or booleans."
                            )
                    rows.append(row)
                result.append(ExcelSheet(sheet.title, rows))
            if not result:
                raise ValueError("Workbook has no worksheets.")
            return result
        finally:
            workbook.close()
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError(
            "Unable to read this .xlsx file. Save it as a standard Excel workbook."
        ) from exc


def columns(sheet: ExcelSheet, header_row: int) -> dict[str, str]:
    if not 1 <= header_row <= min(len(sheet.rows), 100):
        raise ValueError("Header row must identify a row between 1 and 100 in the selected sheet.")
    return {
        get_column_letter(i + 1): f"{get_column_letter(i + 1)} — {label or '(blank header)'}"
        for i, label in enumerate(sheet.rows[header_row - 1])
    }


def convert_excel(
    sheet: ExcelSheet,
    header_row: int,
    mapping: dict[str, str],
    *,
    name: str,
    codebook_id: str,
    version: str,
) -> CodebookSpec:
    available = columns(sheet, header_row)
    if not name.strip() or not codebook_id.strip() or not version.strip():
        raise ValueError("Codebook name, ID, and version are required.")
    if set(mapping) - set(FIELDS):
        raise ValueError("Unknown codebook field in column mapping.")
    selected = {k: v for k, v in mapping.items() if v}
    if not REQUIRED <= selected.keys():
        raise ValueError("Map Code ID, Code label, and Definition before previewing.")
    if len(set(selected.values())) != len(selected):
        raise ValueError("Choose a different column for each mapped field.")
    if set(selected.values()) - available.keys():
        raise ValueError("A mapped column is not present in this sheet.")
    indices = {key: list(available).index(value) for key, value in selected.items()}
    codes = []
    for row_number, row in enumerate(sheet.rows[header_row:], header_row + 1):
        if not any(row):
            continue
        values = {key: row[index] if index < len(row) else "" for key, index in indices.items()}
        missing = [FIELDS[key] for key in REQUIRED if not values.get(key)]
        if missing:
            raise ValueError(
                f"Row {row_number}: missing {', '.join(sorted(missing))}. "
                "Remove notes or incomplete rows from the selected table."
            )
        values["parent_id"] = values.get("parent_id") or None
        values["examples"] = [
            line.strip() for line in values.get("examples", "").splitlines() if line.strip()
        ]
        codes.append(CodebookEntry(**values))
    if not codes:
        raise ValueError("No code rows found below the selected header.")
    return CodebookSpec(
        id=codebook_id.strip(), name=name.strip(), version=version.strip(), codes=codes
    )
