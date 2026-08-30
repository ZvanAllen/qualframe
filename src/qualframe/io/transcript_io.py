from __future__ import annotations

import re
from pathlib import Path


def participant_id(path: str | Path) -> str:
    """Extract the participant ID prefix from a transcript filename.

    E.g. ``'data/raw/transcripts/ON AMN/1A_Transcript.docx'`` → ``'1A'``
    """
    stem = Path(path).stem
    m = re.match(r'^([A-Za-z0-9]+)', stem)
    return m.group(1) if m else stem


def transcript_to_text(path: Path) -> str:
    """Convert a transcript file to plain text.

    Supported formats:
    - .docx  — Word document (paragraphs joined by newlines)
    - .txt   — UTF-8 plain text (universal newlines normalized to LF)

    Evidence offsets address the returned text, not bytes in the original file.
    """
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _docx_to_text(path)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported transcript format: {suffix!r} ({path.name})")


def codebook_xlsx_to_text(path: Path) -> str:
    """Serialize a CFIR codebook .xlsx file to plain text.

    The expected layout is a single sheet where:
    - Rows with only col A populated are domain headers or preamble text.
    - Rows with both col A and col B are (Construct Name, Construct Definition) pairs.
    """
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            a = str(row[0]).strip() if row[0] is not None else ""
            b = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            if not a:
                continue
            if b and b.lower() != "construct definition":
                # construct row
                lines.append(f"{a}: {b}")
            else:
                # header / preamble block — preserve as-is
                lines.append(a)
            lines.append("")
    return "\n".join(lines).strip()


def _docx_to_text(path: Path) -> str:
    """Convert a .docx file to plain text, inserting [Page N] markers at page breaks.

    Detects both explicit page breaks (<w:br w:type="page"/>) and Word's rendered
    page break markers (<w:lastRenderedPageBreak/>), which are the most reliable
    indicators of page boundaries in interview transcripts.
    """
    from docx import Document  # python-docx

    _W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    doc = Document(str(path))
    page = 1
    lines: list[str] = [f"[Page {page}]"]

    for para in doc.paragraphs:
        el = para._element
        has_page_break = (
            el.find(f".//{{{_W}}}br[@{{{_W}}}type='page']") is not None
            or el.find(f".//{{{_W}}}lastRenderedPageBreak") is not None
        )
        if has_page_break:
            page += 1
            lines.append(f"[Page {page}]")
        if para.text.strip():
            lines.append(para.text)

    return "\n\n".join(lines)


def discover_transcripts(transcripts_dir: Path) -> dict[str, list[Path]]:
    """Discover transcript files grouped by network subfolder.

    Expected layout::

        transcripts_dir/
          ON AMN/participant_7A.docx
          BC AMN/participant_3B.docx
          ...

    Returns a dict mapping network label → sorted list of transcript Paths.
    Files directly under ``transcripts_dir`` (no subfolder) are grouped under
    the key ``"_ungrouped"``.
    """
    if not transcripts_dir.exists():
        return {}

    result: dict[str, list[Path]] = {}
    extensions = {".docx", ".txt"}

    for item in sorted(transcripts_dir.iterdir()):
        if item.is_dir():
            network = item.name
            files = sorted(f for f in item.iterdir() if f.suffix.lower() in extensions)
            if files:
                result[network] = files
        elif item.is_file() and item.suffix.lower() in extensions:
            result.setdefault("_ungrouped", []).append(item)

    return result
