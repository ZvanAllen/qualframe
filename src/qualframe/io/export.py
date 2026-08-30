from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from qualframe.io.response_parser import parse_markdown_tables
from qualframe.governance import assert_export_allowed


def parsed_rows_to_workbook(rows_by_sheet: dict[str, list[dict[str, str]]], output_path: Path) -> None:
    """Write a dict of sheet_name → row_dicts to an .xlsx file."""
    wb = Workbook()
    wb.remove(wb.active)  # remove default empty sheet

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    wrap = Alignment(wrap_text=True, vertical="top")

    for sheet_name, rows in rows_by_sheet.items():
        if not rows:
            continue
        ws = wb.create_sheet(title=sheet_name[:31])  # Excel sheet name limit
        headers = list(rows[0].keys())

        # Write header row
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap

        # Write data rows
        for row_idx, row in enumerate(rows, start=2):
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row.get(header, ""))
                cell.alignment = wrap

        # Auto-size columns (capped at 60 chars wide)
        for col_idx, header in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = max(
                (len(str(rows[r].get(header, ""))) for r in range(len(rows))),
                default=0,
            )
            ws.column_dimensions[col_letter].width = min(max(len(header), max_len, 10), 60)

        ws.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def export_run(
    run_id: str,
    artifact_root: Path,
    exports_dir: Path,
    require_released: bool = False,
    require_validated: bool = False,
) -> list[Path]:
    """Export all parsed stage artifacts for a run to .xlsx files.

    Returns the list of written file paths.
    """
    manifest_path = artifact_root / "run_manifest.json"
    if not manifest_path.exists():
        raise FileNotFoundError(f"No run_manifest.json found at {artifact_root}")
    assert_export_allowed(artifact_root, require_released)

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if require_validated and manifest.get("status") == "blocked_validation":
        raise PermissionError(
            "Export blocked: one or more artifacts failed structured-output validation."
        )
    run_exports_dir = exports_dir / run_id
    run_exports_dir.mkdir(parents=True, exist_ok=True)

    written: list[Path] = []

    for rel_path in manifest.get("artifact_paths", []):
        artifact_path = artifact_root / rel_path
        if not artifact_path.exists():
            continue

        artifact = json.loads(artifact_path.read_text(encoding="utf-8"))
        parsed_output = artifact.get("parsed_output")
        if not parsed_output:
            # Fall back to re-parsing the raw response text (handles artifacts
            # stored before the parser was fixed or when parsing failed at run time)
            response_text = artifact.get("response_text", "")
            if response_text:
                parsed_output = parse_markdown_tables(response_text) or None
        if not parsed_output:
            continue

        stem = artifact_path.stem  # e.g. "2A_CFIR_AllDomains_ChatV5.2Pro_artifact"
        stage_id = artifact.get("stage_id", "unknown")
        out_name = f"{stage_id}__{stem.removesuffix('_artifact')}.xlsx"
        out_path = run_exports_dir / out_name
        parsed_rows_to_workbook(parsed_output, out_path)
        written.append(out_path)

    return written
