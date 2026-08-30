from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from qualframe.models import WorkbookValidationResult


REQUIRED_COLUMNS = ["CFIR Domain", "CFIR Construct", "Quote", "Explanation"]


def discover_workbooks(path: Path) -> list[Path]:
    if not path.exists():
        return []
    return sorted(p for p in path.rglob("*.xlsx") if not p.name.startswith("~$"))


def workbook_sheet_names(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    return list(workbook.sheetnames)


def workbook_headers(path: Path) -> dict[str, list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    headers: dict[str, list[str]] = {}
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers[sheet_name] = [str(cell).strip() for cell in first_row if cell is not None]
    return headers


def validate_required_columns(path: Path, required: list[str] | None = None) -> dict[str, list[str]]:
    required = required or REQUIRED_COLUMNS
    issues: dict[str, list[str]] = {}
    for sheet_name, headers in workbook_headers(path).items():
        missing = [column for column in required if column not in headers]
        if missing:
            issues[sheet_name] = missing
    return issues


def _normalize(name: str) -> str:
    """Lowercase and strip common suffixes for fuzzy sheet-name matching."""
    return name.lower().removesuffix(" domain").strip()


def match_domain_sheets(sheets: list[str], domains: list[str]) -> dict[str, str | None]:
    """Map each framework domain name to its actual sheet name, or None if absent."""
    normalized_to_actual = {_normalize(s): s for s in sheets}
    return {domain: normalized_to_actual.get(_normalize(domain)) for domain in domains}


def workbook_to_text(path: Path) -> str:
    """Serialize all sheets as markdown tables for inclusion in an LLM prompt."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        parts.append(f"## Sheet: {sheet_name}\n")
        header = [str(c) if c is not None else "" for c in rows[0]]
        parts.append("| " + " | ".join(header) + " |")
        parts.append("| " + " | ".join("---" for _ in header) + " |")
        for row in rows[1:]:
            cells = [str(c).replace("|", "\\|") if c is not None else "" for c in row]
            parts.append("| " + " | ".join(cells) + " |")
        parts.append("")
    return "\n".join(parts)


def validate_workbook(
    path: Path,
    required_columns: list[str],
    expected_domains: list[str],
) -> WorkbookValidationResult:
    """Validate sheet presence (by domain) and required columns for each matched sheet."""
    sheets = workbook_sheet_names(path)
    domain_map = match_domain_sheets(sheets, expected_domains)
    missing_domains = [d for d, s in domain_map.items() if s is None]

    headers = workbook_headers(path)
    column_issues: dict[str, list[str]] = {}
    for domain, sheet in domain_map.items():
        if sheet is None:
            continue
        missing_cols = [c for c in required_columns if c not in headers.get(sheet, [])]
        if missing_cols:
            column_issues[sheet] = missing_cols

    return WorkbookValidationResult(
        path=str(path),
        valid=not missing_domains and not column_issues,
        missing_domain_sheets=missing_domains,
        column_issues=column_issues,
    )
